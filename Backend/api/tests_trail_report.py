"""The audit trail report.

What has to hold: only the offices in the trail, QMS positions and - with
a reason - the system admin may generate it; every page carries the
request's status and a serial; generating is logged; no names printed.
"""

import hashlib
import io
import zipfile
from unittest import mock

from docx import Document

from api.models import (
    Concurrence, CustomUser, Office, Position, Proposal, TrailReport,
)

from .tests_concurrence import PASSWORD, ConcurrenceFixture


def all_xml(data):
    """Every XML part of the .docx, joined - body, header and footer."""
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return ''.join(
            package.read(name).decode('utf-8') for name in package.namelist()
            if name.endswith('.xml')
        )


def header_text(data):
    header = Document(io.BytesIO(data)).sections[0].header
    return '\n'.join(p.text for p in header.paragraphs)


class TrailFixture(ConcurrenceFixture):

    def setUp(self):
        super().setUp()
        self.proposal_id = self.a_draft()
        self.submit(self.proposal_id)

    def generate(self, user, reason=None, proposal_id=None):
        body = {} if reason is None else {'reason': reason}
        return self.as_(user).post(
            f'/api/proposals/{proposal_id or self.proposal_id}/trail-report/',
            body, format='json',
        )


class WhoMayGenerate(TrailFixture):

    def test_the_requesting_office(self):
        for user in (self.acc_enc, self.acc_head):
            response = self.generate(user)
            self.assertEqual(response.status_code, 200, user.username)
            self.assertTrue(response['Content-Disposition'].endswith('.docx"'))

    def test_a_concurring_office(self):
        self.assertEqual(self.generate(self.bud_enc).status_code, 200)

    def test_an_approving_office(self):
        vp = self.person('vp_head', self.vpaf, Position.HEAD)
        self.assertEqual(self.generate(vp).status_code, 200)

    def test_an_office_outside_the_trail_may_not(self):
        library = Office.objects.create(name='Library', abbreviation='LIB')
        outsider = self.person('lib_head', library, Position.HEAD)
        response = self.generate(outsider)
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.data['reason'], 'not_in_trail')
        self.assertFalse(TrailReport.objects.exists())

    def test_the_imr_and_custodian_may_for_any_request(self):
        qms = Office.objects.create(name='QMS Office', abbreviation='QMS')
        for kind, username in ((Position.IMR, 'imr'),
                               (Position.DOCUMENT_CUSTODIAN, 'dc')):
            user = self.person(username, qms, kind)
            response = self.generate(user)
            self.assertEqual(response.status_code, 200, username)
        titles = set(TrailReport.objects.values_list('generated_as', flat=True))
        self.assertEqual(len(titles), 2)
        self.assertTrue(all(t.startswith('QMS Office — ') for t in titles))

    def test_a_system_admin_must_give_a_reason(self):
        admin = CustomUser.objects.create_user(
            username='sysadmin', password=PASSWORD, is_approved=True,
            system_role=CustomUser.SYSTEM_ADMIN,
        )
        response = self.generate(admin)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data['reason'], 'reason_required')

        response = self.generate(admin, reason='Requested by the external auditor.')
        self.assertEqual(response.status_code, 200)
        report = TrailReport.objects.get()
        self.assertEqual(report.generated_as, 'System administrator')
        self.assertEqual(report.reason, 'Requested by the external auditor.')

    def test_an_admin_with_a_post_in_the_trail_is_not_asked(self):
        self.acc_head.system_role = CustomUser.SYSTEM_ADMIN
        self.acc_head.save(update_fields=['system_role'])
        self.assertEqual(self.generate(self.acc_head).status_code, 200)
        self.assertIn('Accounting Office', TrailReport.objects.get().generated_as)


class WhatItSays(TrailFixture):

    def test_an_open_request_is_stamped_in_progress(self):
        data = self.generate(self.acc_head).content
        self.assertIn('IN PROGRESS — NOT IN EFFECT', header_text(data))

    def test_a_withdrawn_request_says_so_on_every_page(self):
        self.as_(self.acc_head).post(
            f'/api/proposals/{self.proposal_id}/withdraw/',
            {'reason': 'Superseded by the new cash-handling policy.'},
            format='json', HTTP_X_REAUTH_TOKEN=self.token(self.acc_head),
        )
        data = self.generate(self.acc_head).content
        # In the header, which Word repeats on every page.
        self.assertIn('WITHDRAWN — NOT IN EFFECT', header_text(data))
        self.assertIn('Superseded by the new cash-handling policy.', all_xml(data))

    def test_denied_and_effective_stamps(self):
        proposal = Proposal.objects.get(pk=self.proposal_id)
        for status, stamp in ((Proposal.DENIED, 'DENIED BY THE IMR — NOT IN EFFECT'),
                              (Proposal.EFFECTIVE, 'EFFECTIVE')):
            proposal.status = status
            proposal.save(update_fields=['status'])
            self.assertIn(stamp, header_text(self.generate(self.acc_head).content))

    def test_serial_generator_and_page_numbers_on_every_page(self):
        response = self.generate(self.bud_head)
        serial = response['X-Report-Serial']
        xml = all_xml(response.content)
        self.assertIn(serial, header_text(response.content))
        footer = Document(io.BytesIO(response.content)).sections[0].footer
        self.assertIn(serial, footer.paragraphs[0].text)
        self.assertIn('Budget Office — Head', footer.paragraphs[0].text)
        self.assertIn('NUMPAGES', xml)
        self.assertIn(' PAGE ', xml)

    def test_the_form_sections_and_the_trail(self):
        self.decide(self.proposal_id, self.bud_head, Concurrence.RETURN,
                    feedback='Keep the five-day window for large cheques.')
        xml = all_xml(self.generate(self.acc_head).content)
        for text in ('Requested by', 'Department/Unit Head', '2. Concurrence',
                     '3. IMR decision', '4. Approving authority',
                     '5. Document status', '7. Complete trail',
                     'Accounting Office — Encoder', 'Accounting Office — Head',
                     'Keep the five-day window for large cheques.',
                     'does not verify them'):
            self.assertIn(text, xml)

    def test_no_names_are_printed(self):
        for user, name in ((self.acc_enc, 'Juana Encoder-Name'),
                           (self.acc_head, 'Maria Head-Name'),
                           (self.bud_head, 'Jose Budget-Name')):
            user.full_name = name
            user.save(update_fields=['full_name'])
        self.decide(self.proposal_id, self.bud_head, Concurrence.CONCUR)
        xml = all_xml(self.generate(self.acc_head).content)
        for name in ('Juana Encoder-Name', 'Maria Head-Name', 'Jose Budget-Name'):
            self.assertNotIn(name, xml)


class TheLog(TrailFixture):

    def test_each_report_is_logged_with_its_fingerprint(self):
        response = self.generate(self.acc_enc)
        report = TrailReport.objects.get()
        self.assertEqual(report.serial, response['X-Report-Serial'])
        self.assertEqual(report.generated_by, self.acc_enc)
        self.assertEqual(report.generated_as, 'Accounting Office — Encoder')
        self.assertEqual(report.status_at_time, Proposal.CONCURRENCE)
        self.assertEqual(report.sha256, hashlib.sha256(response.content).hexdigest())

    def test_generating_does_not_change_the_requests_own_trail(self):
        proposal = Proposal.objects.get(pk=self.proposal_id)
        before = proposal.events.count()
        self.generate(self.acc_head)
        self.assertEqual(proposal.events.count(), before)

    def test_a_report_that_fails_to_build_is_not_logged(self):
        with mock.patch('api.generation.trail.build', side_effect=RuntimeError):
            with self.assertRaises(RuntimeError):
                self.generate(self.acc_head)
        self.assertFalse(TrailReport.objects.exists())

    def test_the_log_lists_every_report(self):
        self.generate(self.acc_head)
        self.generate(self.bud_head)
        rows = self.as_(self.cmo_head).get(
            f'/api/proposals/{self.proposal_id}/trail-report/log/').data['reports']
        self.assertEqual(len(rows), 2)

    def test_a_serial_can_be_looked_up(self):
        serial = self.generate(self.acc_head)['X-Report-Serial']
        data = self.as_(self.bud_head).get(f'/api/trail-reports/{serial.lower()}/').data
        self.assertEqual(data['serial'], serial)
        self.assertEqual(data['stamp'], 'IN PROGRESS — NOT IN EFFECT')

    def test_an_outsider_cannot_look_one_up(self):
        serial = self.generate(self.acc_head)['X-Report-Serial']
        library = Office.objects.create(name='Library', abbreviation='LIB')
        outsider = self.person('lib_head', library, Position.HEAD)
        response = self.as_(outsider).get(f'/api/trail-reports/{serial}/')
        self.assertEqual(response.status_code, 404)


class TheDetailPageOffersIt(TrailFixture):

    def access(self, user):
        return self.as_(user).get(
            f'/api/proposals/{self.proposal_id}/full/').data['trail_report']

    def test_to_an_office_in_the_trail_without_a_reason(self):
        self.assertEqual(self.access(self.acc_enc),
                         {'can_generate': True, 'needs_reason': False})

    def test_to_a_system_admin_with_a_reason(self):
        admin = CustomUser.objects.create_user(
            username='sysadmin', password=PASSWORD, is_approved=True,
            system_role=CustomUser.SYSTEM_ADMIN,
        )
        self.assertEqual(self.access(admin),
                         {'can_generate': True, 'needs_reason': True})
