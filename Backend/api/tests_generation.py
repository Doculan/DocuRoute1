"""The documents generated at the lock: the DCR, the draft copy, the annex.

What has teeth here:

**No names.** Not printed - position titles only - and not hidden in the
file's properties either. The metadata test uses templates deliberately
dirtied with names, because the real templates are clean now and a test
against them would pass with the scrubbing deleted.

**Nothing on the form moves.** Section 2's new paragraphs must take exactly
the height of the empty lines they replace, or sections 3 to 5 shift - and
at worst onto a second sheet. Word cannot run here, so the invariant is
checked on the file: lines times spacing, plus the spacer, equals the
template's own height. The one-page result itself was confirmed by
rendering through Word during development (PROGRESS.md, 3b).

**Both copies of each signature label.** DrawingML and VML, identically.

Every test that expects an error asserts the reason, not only the status.
"""

import io
import os
import shutil
import tempfile
import zipfile
from unittest import mock

import docx
from django.test import SimpleTestCase
from docx.oxml.ns import qn

from api import documents
from api.document_hygiene import (
    doc_personal_metadata, docx_personal_metadata, scrub_document,
)
from api.generation import common, dcr, pages
from api.models import (
    Attachment, AuditEvent, Concurrence, ManualSeriesOffice, Proposal,
)
from api.tests_concurrence import ConcurrenceFixture

MC = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'

TABLE = (
    "|Responsibility||Activity|\n"
    "| --- | --- | --- |\n"
    "| Cashier | 1. | Releases the cheque within one working day. |\n"
    "| Accounting Staff | 2. | Verifies the request. |"
)
LONG_REASON = '\n'.join([
    'The 2026 management review found that cheques were released between three '
    'and eleven days after approval, and that suppliers were not told which '
    'window applied. This amendment fixes a single release window, names the '
    'office accountable for it, and aligns the procedure table with the policy '
    'so that the two no longer disagree.'
] * 5)


def xml_text(attachment, part='word/document.xml'):
    with zipfile.ZipFile(attachment.file.path) as package:
        return package.read(part).decode('utf-8')


def all_text(attachment):
    """Every text node in every XML part - headers, boxes and properties."""
    out = []
    with zipfile.ZipFile(attachment.file.path) as package:
        for name in package.namelist():
            if name.endswith('.xml'):
                out.append(package.read(name).decode('utf-8', 'replace'))
    return '\n'.join(out)


class PackageFixture(ConcurrenceFixture):

    def setUp(self):
        super().setUp()
        # Names that must never reach a generated document.
        for user, first, last in ((self.acc_enc, 'Marisol', 'Dagami'),
                                  (self.acc_head, 'Teodoro', 'Balangiga'),
                                  (self.bud_head, 'Rosalinda', 'Tanauan'),
                                  (self.cmo_head, 'Evangeline', 'Palo')):
            user.first_name, user.last_name = first, last
            user.save(update_fields=['first_name', 'last_name'])
        self.names = ['Marisol', 'Dagami', 'Teodoro', 'Balangiga', 'Rosalinda',
                      'Tanauan', 'Evangeline', 'Palo', 'acc_enc', 'acc_head',
                      'bud_head', 'cmo_head']

    def locked(self, reason=None, sections=None):
        proposal_id = self.a_draft(sections or [
            (self.s1, "The Cashier shall release the cheque within one working day."),
            (self.s2, TABLE),
        ])
        if reason is not None:
            proposal = Proposal.objects.get(pk=proposal_id)
            version = proposal.current_version()
            version.overall_reason = reason
            version.save(update_fields=['overall_reason'])
        self.submit(proposal_id)
        self.decide(proposal_id, self.bud_head, Concurrence.CONCUR)
        self.decide(proposal_id, self.cmo_head, Concurrence.CONCUR)
        return Proposal.objects.get(pk=proposal_id)

    def attachment(self, proposal, kind):
        return proposal.attachments.get(kind=kind)


# --- the package ---------------------------------------------

class PackageTests(PackageFixture):

    def test_locking_makes_the_documents_and_waits_for_signature(self):
        proposal = self.locked()
        self.assertEqual(proposal.status, Proposal.AWAITING_SIGNATURE)
        self.assertEqual(
            set(proposal.attachments.values_list('kind', flat=True)),
            {Attachment.DCR_GENERATED, Attachment.PAGES_GENERATED,
             Attachment.CONCURRENCE_RECORD},
        )
        for attachment in proposal.attachments.all():
            self.assertTrue(attachment.original_filename.startswith(proposal.dcr_number))
            self.assertEqual(attachment.size_bytes, os.path.getsize(attachment.file.path))
            self.assertEqual(attachment.created_by, self.cmo_head)
        self.assertTrue(proposal.events.filter(event=AuditEvent.DOCUMENTS_GENERATED).exists())

    def test_without_other_offices_there_is_no_annex(self):
        """Nothing to record, and a short reason: two documents, not three."""
        ManualSeriesOffice.objects.filter(series=self.series).exclude(
            office=self.accounting).delete()
        proposal_id = self.a_draft()
        self.submit(proposal_id)
        proposal = Proposal.objects.get(pk=proposal_id)
        self.assertEqual(proposal.status, Proposal.AWAITING_SIGNATURE)
        self.assertEqual(
            set(proposal.attachments.values_list('kind', flat=True)),
            {Attachment.DCR_GENERATED, Attachment.PAGES_GENERATED},
        )

    def test_no_person_is_named_anywhere(self):
        """Position titles only - in the page, the boxes and the properties."""
        proposal = self.locked()
        for attachment in proposal.attachments.all():
            text = all_text(attachment)
            for name in self.names:
                self.assertNotIn(name, text, f'{name} in {attachment.kind}')

    def test_a_changed_template_fails_the_lock_and_leaves_it_workable(self):
        """Loudly, inside the lock - never a misfilled form."""
        proposal_id = self.a_draft()
        self.submit(proposal_id)
        self.decide(proposal_id, self.bud_head, Concurrence.CONCUR)
        with mock.patch.object(dcr, 'DCR_TEMPLATE', common.PAGES_TEMPLATE):
            with self.assertRaises(common.TemplateChanged) as caught:
                self.decide(proposal_id, self.cmo_head, Concurrence.CONCUR)
        self.assertIn('single table', str(caught.exception))
        proposal = Proposal.objects.get(pk=proposal_id)
        self.assertEqual(proposal.status, Proposal.CONCURRENCE)
        self.assertEqual(proposal.attachments.count(), 0)


# --- no personal metadata ------------------------------------

class GeneratedMetadataTests(PackageFixture):
    """The generated files carry no creator or last-modified-by.

    Against templates dirtied with names on purpose: `python-docx` copies
    a template's properties into its output, so this is what a template
    replaced next year with its author still in it would produce.
    """

    def setUp(self):
        super().setUp()
        self.folder = tempfile.mkdtemp(prefix='dirty-templates-')
        self.dirty_dcr = self.dirty(common.DCR_TEMPLATE)
        self.dirty_pages = self.dirty(common.PAGES_TEMPLATE)

    def tearDown(self):
        shutil.rmtree(self.folder, ignore_errors=True)
        super().tearDown()

    def dirty(self, path):
        document = docx.Document(path)
        document.core_properties.author = 'Some Author'
        document.core_properties.last_modified_by = 'Someone Else'
        document.core_properties.comments = 'Drafted by Some Author'
        target = os.path.join(self.folder, os.path.basename(path))
        document.save(target)
        found = docx_personal_metadata(target)
        assert {'dc:creator', 'cp:lastModifiedBy'} <= set(found), found
        return target

    def test_generated_documents_carry_no_creator_or_last_modified_by(self):
        with mock.patch.object(dcr, 'DCR_TEMPLATE', self.dirty_dcr), \
                mock.patch.object(pages, 'PAGES_TEMPLATE', self.dirty_pages):
            proposal = self.locked()

        kinds = set()
        for attachment in proposal.attachments.all():
            kinds.add(attachment.kind)
            properties = docx.Document(attachment.file.path).core_properties
            self.assertEqual(properties.author, '', attachment.kind)
            self.assertEqual(properties.last_modified_by, '', attachment.kind)
            self.assertEqual(docx_personal_metadata(attachment.file.path), {},
                             attachment.kind)
        # The DCR and the draft copy, as asked, and the annex too - it is
        # built from python-docx's own default, whose author is "python-docx".
        self.assertEqual(kinds, {Attachment.DCR_GENERATED, Attachment.PAGES_GENERATED,
                                 Attachment.CONCURRENCE_RECORD})


class PackageCheckTests(SimpleTestCase):

    def test_a_file_with_repeated_parts_is_refused(self):
        """Two property parts: which one a reader believes is up to the reader."""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as package:
            package.writestr('word/document.xml', '<w:document/>')
            package.writestr('docProps/core.xml', '<cp:coreProperties/>')
            package.writestr('docProps/core.xml', '<cp:coreProperties/>')
        with self.assertRaises(common.TemplateChanged) as caught:
            common.check_package(buffer.getvalue(), 'x.docx')
        self.assertIn('repeats docProps/core.xml', str(caught.exception))

    def test_a_file_still_carrying_a_name_is_refused(self):
        document = docx.Document()
        document.core_properties.author = 'Some Author'
        buffer = io.BytesIO()
        document.save(buffer)
        with self.assertRaises(common.TemplateChanged) as caught:
            common.check_package(buffer.getvalue(), 'x.docx')
        self.assertIn('dc:creator', str(caught.exception))


# --- the DCR -------------------------------------------------

class DcrContentTests(PackageFixture):

    def test_the_blanks_are_filled(self):
        proposal = self.locked()
        text = xml_text(self.attachment(proposal, Attachment.DCR_GENERATED))
        for expected in (proposal.dcr_number, common.form_date(proposal.locked_at),
                         'Accounting Office', 'FAM 6.02', ' X ',
                         'Sections amended: 3.0 POLICIES; 4.0 PROCEDURES.',
                         dcr.SEE_DRAFT, proposal.current_version().overall_reason):
            self.assertIn(expected, text)

    def test_both_copies_of_each_label_carry_the_position(self):
        """DrawingML for current Word, VML for older readers - identically."""
        proposal = self.locked()
        document = docx.Document(self.attachment(proposal, Attachment.DCR_GENERATED).file.path)
        seen = {'Choice': [], 'Fallback': []}
        for alternate in document.element.body.iter(MC + 'AlternateContent'):
            for branch in ('Choice', 'Fallback'):
                box = alternate.find(MC + branch).find('.//' + qn('w:txbxContent'))
                if box is not None:
                    seen[branch].append(''.join(t.text or '' for t in box.iter(qn('w:t'))))
        for branch, boxes in seen.items():
            self.assertIn('Accounting Office — EncoderRequested by', boxes, branch)
            self.assertIn('Accounting Office — HeadDepartment/Unit Head', boxes, branch)

    def test_section_two_keeps_the_height_of_the_lines_it_replaced(self):
        """So nothing below it moves, and the form stays on one page."""
        proposal = self.locked()
        template = dcr._Form(docx.Document(common.DCR_TEMPLATE))
        filled = docx.Document(self.attachment(proposal, Attachment.DCR_GENERATED).file.path)

        cell = next(c for row in filled.tables[0].rows for c in row.cells
                    if 'CHANGE(S) REQUESTED' in c.text)
        own = [p for p in cell._tc.iter(qn('w:p')) if p.getparent() is cell._tc]
        texts = [dcr._text(p) for p in own]
        heading = next(i for i, t in enumerate(texts) if 'CHANGE(S) REQUESTED' in t)
        reason = next(i for i, t in enumerate(texts) if 'REASON FOR THE CHANGE' in t)
        signature = next(i for i in range(reason + 1, len(own)) if dcr._has_drawing(own[i]))

        def height(paragraphs):
            total = 0.0
            for p in paragraphs:
                spacing = p.find(qn('w:pPr')).find(qn('w:spacing'))
                lines = 1 + len(p.findall('.//' + qn('w:br')))
                total += lines * int(spacing.get(qn('w:line'))) / 20.0
                total += int(spacing.get(qn('w:after'))) / 20.0
            return total

        for area, (start, end) in ((template.changes_area, (heading + 1, reason)),
                                   (template.reason_area, (reason + 1, signature))):
            expected = template.area_height(area)
            self.assertAlmostEqual(height(own[start:end]), expected, delta=0.1)

    def test_a_reason_too_long_for_the_page_goes_to_the_annex_in_full(self):
        proposal = self.locked(reason=LONG_REASON)
        dcr_text = xml_text(self.attachment(proposal, Attachment.DCR_GENERATED))
        self.assertIn(dcr.REASON_IN_ANNEX, dcr_text)
        self.assertNotIn('found that cheques', dcr_text)

        annex = docx.Document(self.attachment(proposal, Attachment.CONCURRENCE_RECORD).file.path)
        paragraphs = [p.text for p in annex.paragraphs]
        self.assertIn('Reason for the change', paragraphs)
        self.assertEqual(
            [p for p in paragraphs if p.startswith('The 2026 management review')],
            [line for line in LONG_REASON.splitlines()],
        )

    def test_a_long_list_of_sections_is_shortened_and_says_so(self):
        names = ['%d.0 A SECTION WITH A FAIRLY LONG HEADING' % i for i in range(1, 41)]
        plan = dcr._fit_changes(names, 490, 70)
        self.assertRegex(plan['paragraphs'][0], r'; and \d+ other sections\.$')
        self.assertEqual(plan['paragraphs'][-1], dcr.SEE_DRAFT)
        short = dcr._fit_changes(names[:2], 490, 70)
        self.assertEqual(short['paragraphs'][0],
                         'Sections amended: %s; %s.' % tuple(names[:2]))

    def test_a_long_office_name_takes_two_lines_and_the_reason_makes_room(self):
        form = dcr._Form(docx.Document(common.DCR_TEMPLATE))
        short = {label: form.fit_title(label, 'Budget Office — Head')
                 for label in form.boxes}
        long_ = {label: form.fit_title(
                     label, 'Office of the Vice President for Administration '
                            'and Finance — Head')
                 for label in form.boxes}
        self.assertTrue(all(len(f['lines']) == 1 for f in short.values()))
        self.assertTrue(all(len(f['lines']) == 2 for f in long_.values()))
        self.assertLess(form.reason_room(long_), form.reason_room(short))


# --- the draft copy ------------------------------------------

class DraftCopyTests(PackageFixture):

    def test_pipe_tables_keep_their_empty_columns(self):
        """An empty cell at either end must not be eaten.

        "||VERSION NO.|..." opens with one. Stripping every pipe from the
        ends would drop it and shift every column left - the empty middle
        column of "|Responsibility||Activity|" survives either way, which
        is why this test checks the ends.
        """
        [(kind, table)] = pages.blocks(TABLE)
        self.assertEqual(kind, 'table')
        self.assertTrue(table['header'])
        self.assertEqual(table['rows'][0], ['Responsibility', '', 'Activity'])
        self.assertEqual(len(table['rows']), 3)

        [(_, edges)] = pages.blocks('||VERSION NO.|1|\n|Title|x||')
        self.assertEqual(edges['rows'], [['', 'VERSION NO.', '1'], ['Title', 'x', '']])

    def test_the_draft_copy_prints_the_changed_sections_as_agreed(self):
        proposal = self.locked(sections=[
            (self.s2, TABLE),
            (self.s1, "The Cashier shall release the cheque within one working day."),
        ])
        attachment = self.attachment(proposal, Attachment.PAGES_GENERATED)
        document = docx.Document(attachment.file.path)
        paragraphs = [p.text for p in document.paragraphs]
        # Document order, not the order they were edited in.
        self.assertLess(paragraphs.index('3.0 POLICIES'), paragraphs.index('4.0 PROCEDURES'))
        self.assertIn('The Cashier shall release the cheque within one working day.', paragraphs)
        self.assertNotIn('| --- | --- | --- |', '\n'.join(paragraphs))
        [table] = document.tables
        self.assertEqual([c.text for c in table.rows[0].cells],
                         ['Responsibility', '', 'Activity'])

    def test_unchanged_sections_are_left_out(self):
        proposal = self.locked(sections=[
            (self.s1, "The Cashier shall release the cheque within one working day."),
        ])
        document = docx.Document(
            self.attachment(proposal, Attachment.PAGES_GENERATED).file.path)
        self.assertNotIn('4.0 PROCEDURES', [p.text for p in document.paragraphs])

    def test_the_header_counts_the_pages(self):
        """"Page 1 of 4", as the university's documents read."""
        proposal = self.locked()
        with zipfile.ZipFile(self.attachment(proposal, Attachment.PAGES_GENERATED).file.path) as package:
            headers = [package.read(n).decode('utf-8') for n in package.namelist()
                       if n.startswith('word/header')]
        numbered = [h for h in headers if ' PAGE ' in h]
        self.assertTrue(numbered)
        for header in numbered:
            self.assertIn('NUMPAGES', header)
            self.assertIn('> of <', header)


# --- the annex -----------------------------------------------

class AnnexTests(PackageFixture):

    def rows(self, proposal):
        annex = docx.Document(self.attachment(proposal, Attachment.CONCURRENCE_RECORD).file.path)
        [table] = annex.tables
        return [[c.text for c in row.cells] for row in table.rows[1:]]

    def test_each_office_is_recorded_by_position(self):
        proposal = self.locked()
        rows = self.rows(proposal)
        self.assertEqual([r[0] for r in rows], ['Budget Office', 'Cash Management Office'])
        self.assertEqual({r[1] for r in rows}, {'Concurs'})
        self.assertEqual([r[3] for r in rows],
                         ['Budget Office — Head', 'Cash Management Office — Head'])

    def test_a_renamed_office_reads_as_it_was(self):
        proposal_id = self.a_draft()
        self.submit(proposal_id)
        self.budget.name = 'Office of Budget and Planning'
        self.budget.save(update_fields=['name'])
        self.decide(proposal_id, self.bud_head, Concurrence.CONCUR)
        self.decide(proposal_id, self.cmo_head, Concurrence.CONCUR)
        rows = self.rows(Proposal.objects.get(pk=proposal_id))
        self.assertIn('Budget Office', [r[0] for r in rows])
        self.assertNotIn('Office of Budget and Planning', [r[0] for r in rows])

    def test_earlier_returns_are_listed_with_their_feedback(self):
        proposal_id = self.a_draft()
        self.submit(proposal_id)
        self.decide(proposal_id, self.bud_head, Concurrence.RETURN,
                    feedback='Say which working day the window starts from.')
        self.client.put(
            '/api/proposals/%d/sections/%d/' % (proposal_id, self.s1.id),
            {'new_text': 'The Cashier shall release the cheque within two working days.'},
            format='json')
        self.client.post('/api/proposals/%d/sections/%d/check/' % (proposal_id, self.s1.id),
                         {}, format='json')
        self.submit(proposal_id)
        self.decide(proposal_id, self.bud_head, Concurrence.CONCUR)
        self.decide(proposal_id, self.cmo_head, Concurrence.CONCUR)
        proposal = Proposal.objects.get(pk=proposal_id)
        annex = docx.Document(self.attachment(proposal, Attachment.CONCURRENCE_RECORD).file.path)
        lines = [p.text for p in annex.paragraphs]
        self.assertIn('Agreed to version 2 of the proposal.', lines)
        self.assertTrue(any(
            line.startswith('Version 1 — returned by Budget Office')
            and line.endswith('Say which working day the window starts from.')
            for line in lines))


# --- the templates themselves --------------------------------

TEMPLATE_FOLDERS = [
    os.path.join(common.TEMPLATE_DIR, 'templates'),
    os.path.join(common.TEMPLATE_DIR, 'MANUAL_BLANK_FORMAT'),
]


def template_files():
    for folder in TEMPLATE_FOLDERS:
        for name in sorted(os.listdir(folder)):
            yield os.path.join(folder, name)


class TemplateHygieneTests(SimpleTestCase):
    """Every template in the repository, as committed."""

    def test_no_template_carries_a_name_in_its_properties(self):
        checked = 0
        for path in template_files():
            if path.endswith('.doc'):
                self.assertEqual(doc_personal_metadata(path), {}, path)
            elif '.docx' in path:
                self.assertEqual(docx_personal_metadata(path), {}, path)
            else:
                continue
            checked += 1
        self.assertGreaterEqual(checked, 3)

    def test_the_templates_survive_python_docx_without_repeating_a_part(self):
        """What the DCR template failed when it was first converted.

        LibreOffice had declared its properties under a relationship type
        `python-docx` does not recognise, so every save added a second
        `docProps/core.xml`.
        """
        for path in (common.DCR_TEMPLATE, common.PAGES_TEMPLATE):
            document = docx.Document(path)
            scrub_document(document)
            buffer = io.BytesIO()
            document.save(buffer)
            names = zipfile.ZipFile(buffer).namelist()
            self.assertEqual(len(names), len(set(names)), path)

    def test_the_detector_finds_a_name_and_the_scrub_removes_it(self):
        document = docx.Document()
        document.core_properties.author = 'Some Author'
        document.core_properties.last_modified_by = 'Someone Else'
        buffer = io.BytesIO()
        document.save(buffer)
        # python-docx's own default also carries a description ("generated
        # by python-docx"), which the detector rightly reports as well.
        found = docx_personal_metadata(io.BytesIO(buffer.getvalue()))
        self.assertEqual(found['dc:creator'], 'Some Author')
        self.assertEqual(found['cp:lastModifiedBy'], 'Someone Else')
        scrub_document(document)
        clean = io.BytesIO()
        document.save(clean)
        self.assertEqual(docx_personal_metadata(io.BytesIO(clean.getvalue())), {})
