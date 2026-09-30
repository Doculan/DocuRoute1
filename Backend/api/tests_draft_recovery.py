import io
from pathlib import Path
from unittest import mock

from django.core.files.base import ContentFile
from django.core.management import call_command, CommandError
from docx import Document

from api.generation.common import PAGES_TEMPLATE
from api.models import Attachment, AuditEvent, CustomUser, Proposal
from api.package_views import package_payload
from api.tests_generation import PackageFixture


class DraftRecoveryTests(PackageFixture):
    def setUp(self):
        super().setUp()
        self.proposal = self.locked(sections=[(self.s1, 'The Cashier shall retain a copy.')])
        self.original = self.attachment(self.proposal, Attachment.PAGES_GENERATED)
        legacy = Document(PAGES_TEMPLATE)
        legacy.add_paragraph(f'Draft copy attached to {self.proposal.dcr_number} · changed sections only')
        # Blank template can contain empty paragraphs; legacy generation removes them.
        for paragraph in list(legacy.paragraphs):
            if not paragraph.text.strip():
                paragraph._element.getparent().remove(paragraph._element)
        legacy.add_paragraph(self.s1.subtitle)
        legacy.add_paragraph('The Cashier shall retain a copy.')
        buffer = io.BytesIO()
        legacy.save(buffer)
        self.original.file.save('legacy.docx', ContentFile(buffer.getvalue()))
        self.original_bytes = Path(self.original.file.path).read_bytes()
        self.admin_user = CustomUser.objects.create_user(
            username='recovery_admin', password='x', role='admin',
            system_role=CustomUser.SYSTEM_ADMIN, is_approved=True)

    def recover(self, apply=False, **kwargs):
        options = dict(proposal=self.proposal.pk, stdout=io.StringIO())
        if apply:
            options.update(apply=True, actor=self.admin_user.pk, reason='Complete the unsigned draft.',
                           accept_current_baseline=True)
        options.update(kwargs)
        call_command('recover_complete_draft', **options)
        return options['stdout'].getvalue()

    def test_dry_run_leaves_records_and_files_unchanged(self):
        ids = list(Attachment.objects.values_list('pk', flat=True))
        events = AuditEvent.objects.count()
        result = self.recover()
        self.assertIn('DRY RUN', result)
        self.assertEqual(ids, list(Attachment.objects.values_list('pk', flat=True)))
        self.assertEqual(events, AuditEvent.objects.count())
        self.original.refresh_from_db()
        self.assertIsNone(self.original.superseded_at)
        self.assertEqual(Path(self.original.file.path).read_bytes(), self.original_bytes)

    def test_replacement_is_complete_dated_preserved_and_audited(self):
        self.recover(apply=True)
        self.original.refresh_from_db()
        current = self.proposal.attachments.get(kind=Attachment.PAGES_GENERATED, superseded_at=None)
        self.assertEqual(current.supersedes_id, self.original.pk)
        self.assertEqual(current.created_by, self.admin_user)
        self.assertEqual(current.replacement_reason, 'Complete the unsigned draft.')
        self.assertIsNotNone(self.original.superseded_at)
        self.assertEqual(Path(self.original.file.path).read_bytes(), self.original_bytes)
        text = '\n'.join(p.text for p in Document(current.file.path).paragraphs)
        self.assertIn('Reconstructed on', text)
        self.assertIn('not an original lock-time snapshot', text)
        self.assertIn(self.s2.content, text)
        self.assertIn('The Cashier shall retain a copy.', text)
        self.assertEqual(self.proposal.events.filter(detail__startswith='Complete draft recovery:').count(), 1)
        payload = package_payload(self.proposal, self.acc_enc)
        generated = payload['generated']
        self.assertIn(current.pk, [a['id'] for a in generated])
        self.assertNotIn(self.original.pk, [a['id'] for a in generated])
        self.assertEqual([a['id'] for a in payload['previous_generated']], [self.original.pk])
        response = self.as_(self.acc_enc).get(
            f'/api/proposals/{self.proposal.pk}/attachments/{self.original.pk}/download/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(b''.join(response.streaming_content), self.original_bytes)
        response.close()

    def test_any_signed_scan_prevents_replacement(self):
        Attachment.objects.create(proposal=self.proposal, version=self.proposal.current_version(),
                                  kind=Attachment.SIGNED_DCR, created_by=self.acc_enc, file='signed.pdf')
        with self.assertRaisesMessage(CommandError, 'Signed scans exist'):
            self.recover(apply=True)

    def test_downstream_package_is_refused(self):
        self.proposal.status = Proposal.WITH_CUSTODIAN
        self.proposal.save(update_fields=['status'])
        with self.assertRaisesMessage(CommandError, 'Only a proposal awaiting signature'):
            self.recover(apply=True)

    def test_recovery_requires_explicit_baseline_acceptance(self):
        with self.assertRaisesMessage(CommandError, '--apply requires'):
            self.recover(apply=True, accept_current_baseline=False)

    def test_recovery_requires_an_administrator(self):
        with self.assertRaisesMessage(CommandError, 'approved, active system administrator'):
            self.recover(apply=True, actor=self.acc_enc.pk)

    def test_changed_section_baseline_must_still_match(self):
        self.s1.content = 'Unrelated text edited after the request.'
        self.s1.save(update_fields=['content'])
        with self.assertRaisesMessage(CommandError, 'no longer matches its agreed baseline'):
            self.recover(apply=True)

    def test_rerun_does_not_replace_an_already_complete_draft(self):
        self.recover(apply=True)
        count = Attachment.objects.count()
        with self.assertRaisesMessage(CommandError, 'not a legacy partial draft'):
            self.recover(apply=True)
        self.assertEqual(Attachment.objects.count(), count)

    def test_failure_preserves_original_and_removes_new_file(self):
        folder = Path(self.original.file.path).parent
        files = set(folder.iterdir())
        with mock.patch.object(AuditEvent.objects, 'create', side_effect=RuntimeError('Audit unavailable')):
            with self.assertRaisesMessage(RuntimeError, 'Audit unavailable'):
                self.recover(apply=True)
        self.original.refresh_from_db()
        self.assertIsNone(self.original.superseded_at)
        self.assertEqual(set(folder.iterdir()), files)
        self.assertEqual(Path(self.original.file.path).read_bytes(), self.original_bytes)
