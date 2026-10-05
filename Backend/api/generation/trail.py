"""The audit trail of one request, for official reporting.

Laid out on the sections of F-QMS-001, so it reads as a record of the
form the university already uses rather than a screenshot of a system:
who requested it, the unit head, concurrence, the IMR, the approving
authority, the document status. The complete chronological trail follows.

**Every page carries its own status**, in the header: a denied or
withdrawn request's trail found in a folder a year later must not read as
a change that went through. It also carries the serial, the time it was
generated and by which position, and page X of Y.

Position titles only, never names - the same rule as every generated
document. The serial leads back to the `TrailReport` row, which holds who
generated it.

Signatures are on paper. The system records uploaded scans and does not
verify them, and the report says so rather than calling anything signed.
"""

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from django.utils import timezone

from ..document_hygiene import scrub_document
from ..models import (
    Attachment, AuditEvent, Concurrence, Proposal, ProposalParticipant,
    QmsDecision,
)
from .common import _PPR_ORDER, check_package, form_date, position_title, put


# What every page says the request came to. Only one of these means the
# change is in force, and the others say plainly that it is not.
STAMPS = {
    Proposal.EFFECTIVE: 'EFFECTIVE',
    Proposal.DENIED: 'DENIED BY THE IMR — NOT IN EFFECT',
    Proposal.WITHDRAWN: 'WITHDRAWN — NOT IN EFFECT',
}
IN_PROGRESS = 'IN PROGRESS — NOT IN EFFECT'

NOT_VERIFIED = (
    'Signatures are on the paper originals held by the Document Custodian. '
    'This system records the scans uploaded and does not verify them.'
)

GREEN = RGBColor(0x1C, 0x7A, 0x5A)
RED = RGBColor(0xA8, 0x32, 0x32)


def stamp_for(proposal):
    return STAMPS.get(proposal.status, IN_PROGRESS)


def build(proposal, serial, generated_at, generated_as):
    """The report as .docx bytes. Writes nothing to the database."""
    version = proposal.current_version()
    document = Document()
    _page_setup(document)
    _header(document, proposal, serial)
    _footer(document, serial, generated_at, generated_as)

    events = list(
        proposal.events.select_related('position').order_by('at', 'id')
    )

    _request(document, proposal, version, events)
    _concurrence(document, proposal, version)
    _imr(document, proposal)
    _approving(document, proposal)
    _document_status(document, proposal)
    _scans(document, proposal)
    _trail(document, events)

    scrub_document(document)
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    check_package(data, filename(proposal, serial))
    return data


def filename(proposal, serial):
    label = proposal.dcr_number or f'Proposal {proposal.id}'
    return f'{label} Audit trail {serial}.docx'


# ─── Page furniture ──────────────────────────────────────────

def _page_setup(document):
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ('left_margin', 'right_margin'):
        setattr(section, side, Mm(20))
    section.top_margin = Mm(32)
    section.bottom_margin = Mm(22)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)
    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(3)


def _header(document, proposal, serial):
    header = document.sections[0].header
    first = header.paragraphs[0]
    title = first.add_run('DOCUMENT CHANGE AUDIT TRAIL')
    title.bold = True
    title.font.size = Pt(10)
    first.add_run(f'    Serial {serial}').font.size = Pt(8.5)

    stamp = stamp_for(proposal)
    line = header.add_paragraph()
    run = line.add_run(stamp)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = GREEN if proposal.status == Proposal.EFFECTIVE else RED
    _box(line)

    ident = header.add_paragraph()
    manual = proposal.manual
    dcr = proposal.dcr_number or 'DCR not numbered'
    ident.add_run(
        f'Proposal {proposal.id} · {dcr} · {manual.title}'
    ).font.size = Pt(8.5)


def _footer(document, serial, generated_at, generated_as):
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    when = timezone.localtime(generated_at)
    text = (f'Generated {form_date(generated_at)}, {clock(when)} by '
            f'{generated_as} · Serial {serial} · Page ')
    _small(paragraph.add_run(text))
    _field(paragraph, 'PAGE')
    _small(paragraph.add_run(' of '))
    _field(paragraph, 'NUMPAGES')


def _small(run):
    run.font.size = Pt(7.5)
    return run


def _field(paragraph, instruction):
    """A Word field - PAGE or NUMPAGES - which Word fills in when it opens."""
    def run_with(child):
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        size = OxmlElement('w:sz')
        size.set(qn('w:val'), '15')
        props.append(size)
        run.append(props)
        run.append(child)
        paragraph._p.append(run)

    def char(kind):
        element = OxmlElement('w:fldChar')
        element.set(qn('w:fldCharType'), kind)
        return element

    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = f' {instruction} '
    placeholder = OxmlElement('w:t')
    placeholder.text = '1'
    for child in (char('begin'), instr, char('separate'), placeholder, char('end')):
        run_with(child)


def _box(paragraph):
    """A ruled box round the status, so it reads as a stamp."""
    props = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        edge = OxmlElement(f'w:{side}')
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '12')
        edge.set(qn('w:space'), '3')
        edge.set(qn('w:color'), 'auto')
        borders.append(edge)
    put(props, borders, _PPR_ORDER)


# ─── Body helpers ────────────────────────────────────────────

def _heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(10.5)


def _table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ''
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
    _row_rules(table.rows[0], header=True)
    for row in rows:
        added = table.add_row()
        _row_rules(added)
        for cell, text in zip(added.cells, row):
            cell.text = text or '—'
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Mm(width)
    return table


def _row_rules(row, header=False):
    """Never split a row across pages; repeat the heading row on each."""
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement('w:cantSplit'))
    if header:
        props.append(OxmlElement('w:tblHeader'))


def clock(moment):
    """'5:56 PM', not '05:56 PM'."""
    return f'{moment:%I:%M %p}'.lstrip('0')


def _fields(document, pairs):
    """Two-column label/value table, the form's own shape."""
    table = document.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for label, value in pairs:
        cells = table.add_row().cells
        cells[0].text = ''
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].text = value or '—'
        cells[0].width, cells[1].width = Mm(55), Mm(115)
    return table


def _note(document, text):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)


def _event_title(event):
    if event is None:
        return None
    if event.position_id:
        return position_title(event.office_name_at_time, event.position.kind)
    return event.office_name_at_time


def _last(events, kind):
    matching = [e for e in events if e.event == kind]
    return matching[-1] if matching else None


# ─── Sections ────────────────────────────────────────────────

def _request(document, proposal, version, events):
    """Sections 1 and 2 of the form: the request and its reason."""
    _heading(document, '1. Request (F-QMS-001 sections 1–2)')
    created = _last(events, AuditEvent.CREATED)
    submitted = _last(events, AuditEvent.SUBMITTED)
    manual = proposal.manual
    office = created.office_name_at_time if created else proposal.initiating_office.name

    pairs = [
        ('Document', (f'{manual.series.title} — ' if manual.series_id else '') + manual.title),
        ('Requesting office', office),
        ('DCR number', proposal.dcr_number or 'Not numbered — never locked'),
        ('Status when generated', proposal.get_status_display()),
        ('Requested by', _event_title(created) or 'Not recorded'),
        ('Date requested', form_date(proposal.created_at)),
        ('Department/Unit Head',
         f'{_event_title(submitted)}, submitted {form_date(submitted.at)}'
         if submitted else 'Not submitted'),
        ('Version', str(version.number) if version else '—'),
        ('Reason for change', (version.overall_reason or '').strip() if version else ''),
    ]
    if proposal.status == Proposal.WITHDRAWN:
        withdrawn = _last(events, AuditEvent.WITHDRAWN)
        pairs.append((
            'Withdrawn',
            f'{form_date(proposal.withdrawn_at)} by {_event_title(withdrawn)}: '
            f'{proposal.withdrawn_reason}' if withdrawn else form_date(proposal.withdrawn_at),
        ))
    _fields(document, pairs)


def _concurrence(document, proposal, version):
    _heading(document, '2. Concurrence')
    participants = list(
        proposal.participants.filter(role=ProposalParticipant.CONCURRING)
        .order_by('office_name_at_time')
    )
    if not proposal.participants.exists():
        _note(document, 'Not reached: the request was never submitted.')
        return
    if not participants:
        _note(document, 'No other office concurs on this document.')
        return

    decisions = {
        c.office_id: c for c in Concurrence.objects.filter(version=version)
        .select_related('recorded_by_position')
    } if version else {}
    rows = []
    for participant in participants:
        decision = decisions.get(participant.office_id)
        if decision is None:
            pending = ('Not yet decided' if proposal.status == Proposal.CONCURRENCE
                       else 'Not decided')
            rows.append([participant.office_name_at_time, pending, '', ''])
            continue
        title = (
            position_title(participant.office_name_at_time,
                           decision.recorded_by_position.kind)
            if decision.recorded_by_position_id else
            f'{participant.office_name_at_time} — Head'
        )
        rows.append([
            participant.office_name_at_time, decision.get_decision_display(),
            form_date(decision.recorded_at), title,
        ])
    _table(document, ['Office', f'Decision on version {version.number}',
                      'Date', 'Recorded by'], rows, widths=[48, 27, 37, 58])

    names = {p.office_id: p.office_name_at_time for p in participants}
    earlier = Concurrence.objects.filter(
        version__proposal=proposal, decision=Concurrence.RETURN,
    ).select_related('version', 'office').order_by('version__number', 'recorded_at')
    if earlier:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.add_run('Returns').bold = True
        for returned in earlier:
            office = names.get(returned.office_id, returned.office.name)
            line = (f'Version {returned.version.number} returned by {office}, '
                    f'{form_date(returned.recorded_at)}')
            if returned.feedback.strip():
                line += f': {returned.feedback.strip()}'
            document.add_paragraph(line, style='List Bullet')


def _imr(document, proposal):
    _heading(document, '3. IMR decision (F-QMS-001 section 3)')
    decisions = list(
        proposal.qms_decisions.filter(stage=QmsDecision.IMR).order_by('decided_at')
    )
    if not decisions:
        _note(document, 'Not yet reached.')
        return
    for decision in decisions:
        pairs = [
            ('Decision', decision.get_outcome_display()),
            ('Decided by', decision.decided_as),
            ('Date', form_date(decision.decided_at)),
        ]
        if decision.comments.strip():
            pairs.append(('Comments', decision.comments.strip()))
        _fields(document, pairs)


def _approving(document, proposal):
    _heading(document, '4. Approving authority (F-QMS-001 section 4)')
    route = list(
        proposal.participants.filter(role=ProposalParticipant.APPROVING)
        .order_by('route_order')
    )
    if route:
        document.add_paragraph(
            'Approval route: ' + ' → '.join(p.office_name_at_time for p in route)
        )
    scans = list(
        proposal.attachments.filter(kind=Attachment.APPROVED_DCR,
                                    superseded_at__isnull=True)
        .order_by('created_at')
    )
    if scans:
        for scan in scans:
            document.add_paragraph(
                f'Signed copy recorded {form_date(scan.created_at)} by '
                f'{scan.created_as or "—"}.'
            )
    else:
        _note(document, 'No signed copy recorded.')


def _document_status(document, proposal):
    _heading(document, '5. Document status (F-QMS-001 section 5)')
    returns = list(
        proposal.qms_decisions.filter(
            stage=QmsDecision.CUSTODIAN, outcome=QmsDecision.RETURN,
        ).order_by('decided_at')
    )
    for returned in returns:
        line = (f'Returned for package defects {form_date(returned.decided_at)} '
                f'by {returned.decided_as}')
        if returned.comments.strip():
            line += f': {returned.comments.strip()}'
        document.add_paragraph(line, style='List Bullet')

    status = getattr(proposal, 'document_status', None)
    if status is None:
        _note(document, 'Not yet recorded.')
        return
    pairs = [
        ('Document number', status.document_number),
        ('Version', status.version),
        ('Revision', status.revision),
        ('Effectivity date', form_date_date(status.effective_on)),
        ('Recorded by', status.recorded_as),
        ('Recorded on', form_date(status.recorded_at)),
    ]
    if status.updated_in_ids_on:
        pairs.insert(4, ('Updated in IDS', form_date_date(status.updated_in_ids_on)))
    if status.superseded_at:
        pairs.append(('Superseded', f'{form_date(status.superseded_at)}: '
                                    f'{status.correction_reason}'))
    _fields(document, pairs)


def form_date_date(day):
    return f'{day:%B} {day.day}, {day.year}'


def _scans(document, proposal):
    _heading(document, '6. Signed copies on file')
    scans = list(
        proposal.attachments.filter(kind__in=Attachment.SCAN_KINDS)
        .order_by('created_at')
    )
    if not scans:
        _note(document, 'None recorded.')
    else:
        rows = [[
            scan.get_kind_display(),
            scan.original_filename,
            form_date(scan.created_at),
            scan.created_as,
            ('Superseded ' + form_date(scan.superseded_at)
             + (f': {scan.replacement_reason}' if scan.replacement_reason else ''))
            if scan.superseded_at else 'Current',
        ] for scan in scans]
        _table(document, ['Copy', 'File', 'Recorded', 'By', 'State'], rows,
               widths=[32, 45, 25, 40, 28])
    _note(document, NOT_VERIFIED)


def _trail(document, events):
    _heading(document, '7. Complete trail')
    rows = []
    for event in events:
        when = timezone.localtime(event.at)
        rows.append([
            f'{form_date(event.at)}, {clock(when)}',
            event.get_event_display(),
            _event_title(event),
            event.detail.strip(),
        ])
    if not rows:
        _note(document, 'No events recorded.')
        return
    _table(document, ['Date and time', 'Event', 'By', 'Detail'], rows,
           widths=[32, 30, 42, 66])
