"""The draft copy: the changed sections, printed on the document template.

The DCR asks for a draft copy of the document to be attached. This is it -
the agreed text of every changed section, in document order, on
`MANUAL_BLANK.docx`.

**One file for the request, not one per section.** Separate files would
each number their pages from one, and a sheaf of "Page 1 of 1" is not a
draft copy of anything. One file makes "Page 2 of 5" mean something.

**The header is left blank** for hand-filling, as decided: the revision
number and effectivity date are only known after approval. The one thing
added to it is the total page count beside the page number, because the
template has "Page 1" and the university's documents read "Page 1 of 4".

**The text is printed as agreed.** No cleaning, no re-flowing: each line
of the section becomes a paragraph exactly as the offices concurred with
it. The only structure recognised is the pipe table the extraction wrote,
because printing its pipes and dashes would not be the document either.
"""

import re
from copy import deepcopy

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..models import Attachment
from .common import _TBLPR_ORDER, PAGES_TEMPLATE, changed_sections, put, save

# The table separator row, `|---|---|` or `|:--|--:|` - the same rule the
# extraction used when it wrote these tables (ml/ocr_engine.py).
_SEPARATOR = re.compile(r'^\|?(?:\s*:?-{2,}:?\s*\|)+\s*:?-{2,}:?\s*\|?$')


def generate(proposal, version, actor):
    document = Document(PAGES_TEMPLATE)
    _add_total_pages(document)

    body = document.element.body
    for paragraph in list(body.findall(qn('w:p'))):
        if not ''.join(t.text or '' for t in paragraph.iter(qn('w:t'))).strip():
            body.remove(paragraph)

    note = document.add_paragraph()
    run = note.add_run(
        f'Draft copy attached to {proposal.dcr_number} · changed sections only'
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    for index, change in enumerate(changed_sections(version)):
        heading = document.add_paragraph()
        heading.add_run(change.section.subtitle.strip()).bold = True
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(12 if index else 6)

        for kind, content in blocks(change.new_text):
            if kind == 'table':
                _add_table(document, content)
            else:
                document.add_paragraph(content)

    return [save(
        document, proposal, version, Attachment.PAGES_GENERATED, actor,
        filename=f'{proposal.dcr_number} Draft copy {proposal.manual.title}.docx',
    )]


def split_row(line):
    """One table row's cells.

    Strips exactly one delimiting pipe from each end, as the extraction
    does. The danger is an empty cell at an end: "||VERSION NO.|..." opens
    with one, and stripping every pipe would eat it and shift every column
    left. (An empty *middle* cell, as in "|Responsibility||Activity|",
    survives either way.)
    """
    s = line.strip()
    if s.startswith('|'):
        s = s[1:]
    if s.endswith('|'):
        s = s[:-1]
    return [cell.strip() for cell in s.split('|')]


def blocks(text):
    """The section as a sequence of ('paragraph', line) and ('table', rows).

    Blank lines separate; they are not printed, since the paragraph style
    already spaces paragraphs apart.
    """
    out, table = [], []

    def flush():
        if table:
            out.append(('table', _table_rows(table)))
            table.clear()

    for line in (text or '').splitlines():
        if line.strip().startswith('|'):
            table.append(line)
            continue
        flush()
        if line.strip():
            out.append(('paragraph', line.rstrip()))
    flush()
    return out


def _table_rows(lines):
    header = len(lines) > 1 and bool(_SEPARATOR.match(lines[1].strip()))
    rows = [split_row(line) for line in lines if not _SEPARATOR.match(line.strip())]
    width = max(len(row) for row in rows)
    rows = [row + [''] * (width - len(row)) for row in rows]
    return {'rows': rows, 'header': header}


def _add_table(document, table):
    rows = table['rows']
    grid = document.add_table(rows=len(rows), cols=len(rows[0]))
    tblPr = grid._tbl.tblPr

    width = OxmlElement('w:tblW')
    width.set(qn('w:w'), '5000')
    width.set(qn('w:type'), 'pct')
    put(tblPr, width, _TBLPR_ORDER)

    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        line = OxmlElement(f'w:{edge}')
        line.set(qn('w:val'), 'single')
        line.set(qn('w:sz'), '4')
        line.set(qn('w:space'), '0')
        line.set(qn('w:color'), '000000')
        borders.append(line)
    put(tblPr, borders, _TBLPR_ORDER)

    for r, values in enumerate(rows):
        header = table['header'] and r == 0
        if header:
            # Repeated at the top of every page the table runs onto.
            trPr = grid.rows[r]._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:tblHeader'))
        for c, value in enumerate(values):
            cell = grid.cell(r, c)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.bold = header

    # A little air before whatever follows the table.
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _add_total_pages(document):
    """"Page 1" becomes "Page 1 of 4", in every header that numbers pages."""
    for rel in document.part.rels.values():
        if rel.reltype != RT.HEADER:
            continue
        root = rel.target_part.element
        for instruction in root.iter(qn('w:instrText')):
            if instruction.text.strip().split()[0] != 'PAGE':
                continue
            instruction_run = instruction.getparent()
            paragraph = instruction_run.getparent()
            runs = paragraph.findall(qn('w:r'))
            start = runs.index(instruction_run)
            end_run = next(
                r for r in runs[start:]
                if (r.find(qn('w:fldChar')) is not None
                    and r.find(qn('w:fldChar')).get(qn('w:fldCharType')) == 'end')
            )
            rPr = instruction_run.find(qn('w:rPr'))

            def run_with(child):
                run = OxmlElement('w:r')
                if rPr is not None:
                    run.append(deepcopy(rPr))
                run.append(child)
                return run

            def field_char(kind):
                element = OxmlElement('w:fldChar')
                element.set(qn('w:fldCharType'), kind)
                return element

            def text(value):
                element = OxmlElement('w:t')
                element.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                element.text = value
                return element

            instr = OxmlElement('w:instrText')
            instr.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            instr.text = ' NUMPAGES \\* MERGEFORMAT '

            new = [
                run_with(text(' of ')),
                run_with(field_char('begin')),
                run_with(instr),
                run_with(field_char('separate')),
                run_with(text('1')),
                run_with(field_char('end')),
            ]
            anchor = end_run
            for element in new:
                anchor.addnext(element)
                anchor = element
            break
