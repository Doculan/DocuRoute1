"""The annex to the DCR: the record of concurrence, and the reason if long.

The form has no concurrence section, so this is the paper trail for a
stage it does not capture: which offices agreed, when, and in which
position. Office names are the ones frozen when the request was
submitted, so a rename next year does not rewrite what this said.

It also carries the reason for the change when the reason is too long to
fit section 2 on one page. Section 2 then says so, and the reason is still
printed in full, verbatim, where the signatories will read it.

Generated only when there is something to put in it.
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt

from api.models import Attachment, Concurrence, Position, ProposalParticipant
from .common import form_date, position_title, save
from .dcr import plan


def generate(proposal, version, actor):
    reason_in_annex = plan(proposal, version)['reason_in_annex']
    concurring = list(
        proposal.participants.filter(role=ProposalParticipant.CONCURRING)
        .order_by('office_name_at_time')
    )
    if not concurring and not reason_in_annex:
        return []

    document = Document()
    _a4(document)
    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)

    title = document.add_paragraph()
    run = title.add_run(f'Annex to Document Change Request {proposal.dcr_number}')
    run.bold = True
    run.font.size = Pt(12)
    document.add_paragraph(
        f'{proposal.manual.title} · from {proposal.initiating_office.name}'
    )

    if reason_in_annex:
        _heading(document, 'Reason for the change')
        for line in version.overall_reason.splitlines():
            if line.strip():
                document.add_paragraph(line.rstrip())

    if concurring:
        _heading(document, 'Record of concurrence')
        names = {p.office_id: p.office_name_at_time for p in proposal.participants.all()}
        decisions = {
            c.office_id: c for c in
            Concurrence.objects.filter(version=version).select_related('recorded_by_position')
        }
        table = document.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for cell, heading in zip(table.rows[0].cells,
                                 ('Office', 'Decision', 'Date', 'Recorded by')):
            cell.paragraphs[0].add_run(heading).bold = True
        for participant in concurring:
            decision = decisions.get(participant.office_id)
            cells = table.add_row().cells
            cells[0].text = participant.office_name_at_time
            if decision is None:
                cells[1].text = '—'
                continue
            cells[1].text = decision.get_decision_display()
            cells[2].text = form_date(decision.recorded_at)
            cells[3].text = _recorder(decision, participant.office_name_at_time)

        document.add_paragraph(
            f'Agreed to version {version.number} of the proposal.'
        ).paragraph_format.space_before = Pt(6)

        earlier = (
            Concurrence.objects.filter(
                version__proposal=proposal, version__number__lt=version.number,
                decision=Concurrence.RETURN,
            ).select_related('version').order_by('version__number', 'recorded_at')
        )
        if earlier:
            _heading(document, 'Earlier versions')
            for returned in earlier:
                office = names.get(returned.office_id, str(returned.office))
                line = (f'Version {returned.version.number} — returned by '
                        f'{office}, {form_date(returned.recorded_at)}')
                if returned.feedback.strip():
                    line += f': {returned.feedback.strip()}'
                document.add_paragraph(line)

    return [save(
        document, proposal, version, Attachment.CONCURRENCE_RECORD, actor,
        filename=f'{proposal.dcr_number} Annex.docx',
    )]


def _recorder(decision, office_name):
    """The position that recorded it - a Head, by the rules - never the person."""
    position = decision.recorded_by_position
    kind = position.kind if position is not None else Position.HEAD
    return position_title(office_name, kind)


def _heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text).bold = True


def _a4(document):
    """The DCR's paper, since this travels stapled to it."""
    from docx.shared import Mm
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(section, side, Mm(25))
